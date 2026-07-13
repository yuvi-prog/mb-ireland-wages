import logging
from datetime import date, datetime

from openpyxl import load_workbook

log = logging.getLogger(__name__)

LOCATION_SHEETS   = ['Blanchardstown', 'Cork', 'Liffey Valley', 'Nutgrove', 'Whitewater']
INCOME_ROW        = 5
DEFAULT_RATE_WDAY = 14.15
DEFAULT_RATE_SUN  = 17.98

# Column indices for the formula/summary columns
COL_NAME     = 3   # C
COL_RATE_MON = 4   # D
COL_RATE_SUN = 5   # E
COL_TOT_WDAY = 18  # R  — total weekday hours
COL_TOT_WDAY_AMT = 19  # S — weekday pay
COL_TOT_SUN  = 20  # T  — total sunday hours
COL_TOT_SUN_AMT  = 21  # U — sunday pay
COL_FINAL    = 24  # X  — final payment


def _norm(name: str) -> str:
    return name.strip().lower()


def _find_week_columns(ws, target_sunday: date):
    for cell in ws[3]:
        if isinstance(cell.value, datetime) and cell.value.date() == target_sunday:
            sunday_col  = cell.column
            weekday_col = sunday_col - 1
            log.info(f"[{ws.title}] Sunday {target_sunday} → col {sunday_col}, Mon-Sat → col {weekday_col}")
            return weekday_col, sunday_col
    log.warning(f"[{ws.title}] No column found for Sunday {target_sunday}")
    return None, None


def _get_sheet_names(ws) -> list:
    names = []
    for row in ws.iter_rows(min_row=7, min_col=COL_NAME, max_col=COL_NAME):
        cell = row[0]
        if cell.value and isinstance(cell.value, str) and cell.value.strip():
            names.append(cell.value.strip())
    return names


def _find_last_staff_row(ws) -> int:
    last = 7
    for row in ws.iter_rows(min_row=7, max_row=50, min_col=COL_NAME, max_col=COL_NAME):
        if row[0].value and isinstance(row[0].value, str) and row[0].value.strip():
            last = row[0].row
    return last


def _find_staff_row(ws, name: str):
    target = _norm(name)
    for row in ws.iter_rows(min_row=7, min_col=COL_NAME, max_col=COL_NAME):
        cell = row[0]
        if cell.value and _norm(str(cell.value)) == target:
            return cell.row
    return None


def _match_name(square_name: str, sheet_names: list):
    sq_norm   = _norm(square_name)
    sq_parts  = sq_norm.split()
    sheet_map = {_norm(n): n for n in sheet_names}

    if sq_norm in sheet_map:
        return sheet_map[sq_norm]
    if sq_parts and sq_parts[0] in sheet_map:
        return sheet_map[sq_parts[0]]
    if len(sq_parts) > 1 and sq_parts[-1] in sheet_map:
        return sheet_map[sq_parts[-1]]
    return None


def _add_new_staff_row(ws, name: str, new_row: int):
    """
    Add a new staff member at new_row with default rates and all required formulas.
    Mirrors the formula pattern of existing staff rows.
    """
    r = new_row
    ws.cell(r, COL_NAME).value     = name
    ws.cell(r, COL_RATE_MON).value = DEFAULT_RATE_WDAY
    ws.cell(r, COL_RATE_SUN).value = DEFAULT_RATE_SUN

    # Weekday hours total: F+H+J+L+N+P
    ws.cell(r, COL_TOT_WDAY).value     = f'=F{r}+H{r}+J{r}+L{r}+N{r}+P{r}'
    # Weekday pay: total hours × Mon-Sat rate
    ws.cell(r, COL_TOT_WDAY_AMT).value = f'=R{r}*$D{r}'
    # Sunday hours total: G+I+K+M+O
    ws.cell(r, COL_TOT_SUN).value      = f'=G{r}+I{r}+K{r}+M{r}+O{r}'
    # Sunday pay: total hours × Sunday rate
    ws.cell(r, COL_TOT_SUN_AMT).value  = f'=T{r}*$E{r}'
    # Final payment: weekday pay + sunday pay + bonus + adjustment
    ws.cell(r, COL_FINAL).value        = f'=S{r}+U{r}+V{r}+W{r}'

    log.info(f"[{ws.title}] Added new staff row for '{name}' at row {r} "
             f"(rates: €{DEFAULT_RATE_WDAY}/h weekday, €{DEFAULT_RATE_SUN}/h Sunday)")


def _clear_week_columns(ws, weekday_col: int, sunday_col: int):
    """Wipe this week's two columns before writing so stale data never persists."""
    for row in ws.iter_rows(min_row=5, min_col=weekday_col, max_col=sunday_col):
        for cell in row:
            if cell.column in (weekday_col, sunday_col):
                if not isinstance(cell.value, str):
                    cell.value = None


def update_excel_wages(
    file_path: str,
    shifts: dict,
    income: dict,
    target_sunday: date = None,
) -> dict:
    if target_sunday is None:
        from square_service import get_previous_week_sunday
        target_sunday = get_previous_week_sunday()

    wb      = load_workbook(file_path)
    summary = {}

    for sheet_name in LOCATION_SHEETS:
        if sheet_name not in wb.sheetnames:
            log.warning(f"Sheet '{sheet_name}' missing — skipping")
            continue

        ws              = wb[sheet_name]
        location_shifts = shifts.get(sheet_name, {})
        location_income = income.get(sheet_name, 0.0)

        weekday_col, sunday_col = _find_week_columns(ws, target_sunday)
        if not weekday_col:
            summary[sheet_name] = {'error': f'Column not found for Sunday {target_sunday}'}
            continue

        # Wipe this week's columns first
        _clear_week_columns(ws, weekday_col, sunday_col)

        # Write income into row 5
        if location_income:
            ws.cell(row=INCOME_ROW, column=weekday_col).value = location_income
            log.info(f"[{ws.title}] Income: €{location_income:.2f} → row {INCOME_ROW}, col {weekday_col}")

        if not location_shifts:
            summary[sheet_name] = {
                'updated': [], 'unmatched': [], 'added': [],
                'note': 'No shifts this week',
                'income': location_income,
            }
            continue

        sheet_names = _get_sheet_names(ws)
        updated   = []
        unmatched = []
        added     = []

        for sq_name, hours in location_shifts.items():
            matched = _match_name(sq_name, sheet_names)

            if not matched:
                # Auto-add new staff member at the bottom
                new_row = _find_last_staff_row(ws) + 1
                _add_new_staff_row(ws, sq_name, new_row)
                sheet_names.append(sq_name)  # keep sheet_names in sync
                matched = sq_name
                added.append(sq_name)

            row = _find_staff_row(ws, matched)
            if not row:
                unmatched.append(sq_name)
                continue

            weekday_h = round(hours['weekday_hours'], 2)
            sunday_h  = round(hours['sunday_hours'],  2)

            if weekday_h > 0:
                ws.cell(row=row, column=weekday_col).value = weekday_h
            if sunday_h > 0:
                ws.cell(row=row, column=sunday_col).value  = sunday_h

            updated.append({'name': matched, 'weekday_hours': weekday_h, 'sunday_hours': sunday_h})
            log.info(f"[{ws.title}] {matched}: {weekday_h}h weekday, {sunday_h}h Sun")

        summary[sheet_name] = {
            'updated':   updated,
            'unmatched': unmatched,
            'added':     added,
            'income':    location_income,
        }
        log.info(f"[{ws.title}] Done — {len(updated)} updated, {len(added)} added, {len(unmatched)} unmatched")

    wb.save(file_path)
    log.info(f"Saved workbook to {file_path}")
    return summary
